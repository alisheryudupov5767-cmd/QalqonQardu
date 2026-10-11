import os, sqlite3, csv, io, uuid
from datetime import datetime
from functools import wraps
from flask import (Flask, g, render_template, request, redirect, url_for,
                   session, flash, send_from_directory, Response, abort)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

BASE = os.path.dirname(os.path.abspath(__file__))
UPLOADS = os.path.join(BASE, "uploads")
DB = os.path.join(BASE, "qalqon.db")
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "o'zgartiring-bu-kalitni")
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
ALLOWED = {"jpg", "jpeg", "png", "webp", "heic", "pdf"}
FACULTIES = ["Fizika-matematika", "Filologiya", "Tarix", "Iqtisodiyot", "Pedagogika",
             "Kimyo-biologiya", "Axborot texnologiyalari", "Jismoniy madaniyat", "Boshqa"]

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB); g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close(_):
    d = g.pop("db", None)
    if d: d.close()

def init_db():
    c = sqlite3.connect(DB)
    c.executescript("""
    CREATE TABLE IF NOT EXISTS admins(id INTEGER PRIMARY KEY, login TEXT UNIQUE, pw TEXT);
    CREATE TABLE IF NOT EXISTS records(
      id INTEGER PRIMARY KEY, kind TEXT, member TEXT, first_name TEXT, last_name TEXT,
      father_name TEXT, faculty TEXT, course TEXT, group_name TEXT, ttj TEXT, tutor TEXT,
      note TEXT, person_type TEXT, violation_type TEXT, evidence_note TEXT,
      doc_photo TEXT, evidence_photo TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS sos(id INTEGER PRIMARY KEY, member TEXT, lat REAL, lng REAL,
      status TEXT DEFAULT 'yangi', created_at TEXT);
    CREATE TABLE IF NOT EXISTS tasks(id INTEGER PRIMARY KEY, text TEXT, urgent INTEGER, created_at TEXT);
    """)
    if not c.execute("SELECT 1 FROM admins").fetchone():
        c.execute("INSERT INTO admins(login,pw) VALUES(?,?)",
                  (os.environ.get("ADMIN_LOGIN", "admin"),
                   generate_password_hash(os.environ.get("ADMIN_PASSWORD", "admin123"))))
    c.commit(); c.close()

def save_file(field):
    f = request.files.get(field)
    if not f or not f.filename: return None
    ext = f.filename.rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED: return None
    name = f"{datetime.now():%Y%m%d%H%M%S}_{uuid.uuid4().hex[:6]}.{ext}"
    f.save(os.path.join(UPLOADS, secure_filename(name)))
    return name

def admin_only(fn):
    @wraps(fn)
    def w(*a, **k):
        if not session.get("admin"): return redirect(url_for("admin_login"))
        return fn(*a, **k)
    return w

@app.route("/")
def home(): return render_template("home.html")

@app.route("/qalqon")
def qalqon():
    tasks = db().execute("SELECT * FROM tasks ORDER BY id DESC LIMIT 10").fetchall()
    return render_template("qalqon.html", tasks=tasks)

@app.route("/qalqon/<kind>", methods=["GET", "POST"])
def form(kind):
    if kind not in ("post", "hudud"): abort(404)
    if request.method == "POST":
        f = request.form
        db().execute("""INSERT INTO records(kind,member,first_name,last_name,father_name,faculty,
          course,group_name,ttj,tutor,note,person_type,violation_type,evidence_note,
          doc_photo,evidence_photo,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
          (kind, f.get("member"), f.get("first_name"), f.get("last_name"), f.get("father_name"),
           f.get("faculty"), f.get("course"), f.get("group_name"), f.get("ttj"), f.get("tutor"),
           f.get("note"), f.get("person_type"), f.get("violation_type"), f.get("evidence_note"),
           save_file("doc_photo"), save_file("evidence_photo"),
           datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
        db().commit()
        flash("Ma'lumot saqlandi va TXTB xodimiga yuborildi.")
        return redirect(url_for("form", kind=kind))
    return render_template("form.html", kind=kind, faculties=FACULTIES)

@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        r = db().execute("SELECT * FROM admins WHERE login=?", (request.form["login"],)).fetchone()
        if r and check_password_hash(r["pw"], request.form["password"]):
            session["admin"] = r["login"]; return redirect(url_for("admin"))
        flash("Login yoki parol noto'g'ri.")
    return render_template("admin_login.html")

@app.route("/admin/logout")
def admin_logout():
    session.clear(); return redirect(url_for("home"))

def query_records():
    a = request.args; where, p = [], []
    for col, key in (("faculty", "faculty"), ("kind", "kind"), ("violation_type", "violation_type")):
        if a.get(key): where.append(f"{col}=?"); p.append(a[key])
    if a.get("date_from"): where.append("date(created_at)>=?"); p.append(a["date_from"])
    if a.get("date_to"): where.append("date(created_at)<=?"); p.append(a["date_to"])
    if a.get("q"):
        where.append("(first_name||' '||last_name||' '||ifnull(group_name,'') LIKE ?)"); p.append(f"%{a['q']}%")
    sql = "SELECT * FROM records" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY id DESC"
    return db().execute(sql, p).fetchall()

@app.route("/admin")
@admin_only
def admin():
    tasks = db().execute("SELECT * FROM tasks ORDER BY id DESC LIMIT 20").fetchall()
    return render_template("admin.html", rows=query_records(), tasks=tasks, faculties=FACULTIES)

@app.route("/admin/table")
@admin_only
def admin_table(): return render_template("_table.html", rows=query_records())

@app.route("/admin/task", methods=["POST"])
@admin_only
def admin_task():
    t = request.form.get("text", "").strip()
    if t:
        db().execute("INSERT INTO tasks(text,urgent,created_at) VALUES(?,?,?)",
                     (t, 1 if request.form.get("urgent") else 0, datetime.now().strftime("%Y-%m-%d %H:%M")))
        db().commit(); flash("Topshiriq yuborildi.")
    return redirect(url_for("admin"))

@app.route("/admin/export")
@admin_only
def export():
    out = io.StringIO(); w = csv.writer(out)
    w.writerow(["ID", "Bo'lim", "Xodim", "Familiya", "Ism", "Otasining ismi", "Fakultet", "Kurs",
                "Guruh", "TTJ", "Tyutor", "Shaxs turi", "Qoidabuzarlik", "Izoh", "Dalil izohi", "Vaqt"])
    for r in query_records():
        w.writerow([r["id"], "Hudud" if r["kind"] == "hudud" else "Post", r["member"], r["last_name"],
                    r["first_name"], r["father_name"], r["faculty"], r["course"], r["group_name"], r["ttj"],
                    r["tutor"], r["person_type"], r["violation_type"], r["note"], r["evidence_note"], r["created_at"]])
    return Response("\ufeff" + out.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=qalqon_hisobot.csv"})

@app.route("/admin/export.docx")
@admin_only
def export_docx():
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.section import WD_ORIENT
    d = Document(); sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(sec, m, Cm(1.2))
    d.styles["Normal"].font.name = "Times New Roman"; d.styles["Normal"].font.size = Pt(10)
    d.add_heading("Qalqon: hodisalar hisoboti", 1)
    d.add_paragraph(f"Qarshi davlat universiteti. Sana: {datetime.now():%d.%m.%Y %H:%M}")
    heads = ["T/r", "Vaqt", "Bo'lim", "F.I.O.", "Fakultet / kurs / guruh", "TTJ", "Tyutor",
             "Shaxs turi", "Qoidabuzarlik", "Izoh", "Qalqon xodimi"]
    t = d.add_table(rows=1, cols=len(heads)); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, heads):
        c.text = ""; c.paragraphs[0].add_run(h).bold = True
    for i, r in enumerate(query_records(), 1):
        vals = [i, r["created_at"], "Hudud" if r["kind"] == "hudud" else "Post",
                f'{r["last_name"] or ""} {r["first_name"] or ""} {r["father_name"] or ""}'.strip(),
                f'{r["faculty"] or ""}, {r["course"] or ""}-kurs, {r["group_name"] or ""}',
                r["ttj"], r["tutor"], r["person_type"], r["violation_type"],
                r["evidence_note"] or r["note"], r["member"]]
        for c, v in zip(t.add_row().cells, vals): c.text = str(v or "")
    buf = io.BytesIO(); d.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": "attachment; filename=qalqon_hisobot.docx"})

@app.route("/uploads/<name>")
@admin_only
def uploads(name): return send_from_directory(UPLOADS, name)

@app.route("/api/sos", methods=["POST"])
def api_sos():
    d = request.get_json(silent=True) or {}
    db().execute("INSERT INTO sos(member,lat,lng,created_at) VALUES(?,?,?,?)",
                 (str(d.get("member") or "Noma'lum")[:80], d.get("lat"), d.get("lng"),
                  datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    db().commit()
    return {"ok": True}

@app.route("/admin/sos.json")
@admin_only
def sos_json():
    rows = db().execute("SELECT * FROM sos WHERE status='yangi' ORDER BY id DESC").fetchall()
    return {"items": [dict(r) for r in rows]}

@app.route("/admin/sos/<int:i>/ack", methods=["POST"])
@admin_only
def sos_ack(i):
    db().execute("UPDATE sos SET status='qabul qilindi' WHERE id=?", (i,)); db().commit()
    return {"ok": True}

init_db()
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
